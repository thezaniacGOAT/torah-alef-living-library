#!/usr/bin/env python3
"""Daily Torah Aleph archive. Existing Dropbox catalog is authoritative.

Auto-caption Word editions and AI summaries are UNVERIFIED drafts.
No public captions => no invented transcript.
"""
from __future__ import annotations
import csv
import datetime as dt
import io
import json
import os
import re
import unicodedata
from html import unescape
from collections import defaultdict
from pathlib import Path

import requests
from docx import Document
from docx.shared import Inches, Pt

HERE = Path(__file__).resolve().parent.parent
SPEAKERS = json.loads((HERE / "speakers.json").read_text(encoding="utf-8"))
ROOT = os.getenv("DROPBOX_ROOT", "/05. TORAH/Torah Aleph/Shiur Library").rstrip("/")
YOUTUBE_KEY = os.getenv("YOUTUBE_API_KEY", "")
OPENAI_KEY = os.getenv("OPENAI_API_KEY", "")
NOW = dt.datetime.now(dt.timezone.utc)
HTTP = requests.Session()
HTTP.headers["User-Agent"] = "TorahAlephArchive/1.0 (public-lecture-transcripts)"


def dropbox_client():
    import dropbox
    required = ("DROPBOX_APP_KEY", "DROPBOX_APP_SECRET", "DROPBOX_REFRESH_TOKEN")
    missing = [k for k in required if not os.getenv(k)]
    if missing:
        raise RuntimeError("Dropbox secrets missing: " + ", ".join(missing))
    return dropbox.Dropbox(app_key=os.environ["DROPBOX_APP_KEY"],
        app_secret=os.environ["DROPBOX_APP_SECRET"],
        oauth2_refresh_token=os.environ["DROPBOX_REFRESH_TOKEN"], timeout=60)


def read_bytes(dbx, path):
    try:
        _, result = dbx.files_download(path)
        return result.content
    except Exception as ex:
        if "not_found" in str(ex).lower():
            return None
        raise


def save(dbx, path, value, overwrite=False):
    """Safety boundary: only the dedicated Shiur Library can be written."""
    import dropbox
    if not path.startswith(ROOT + "/"):
        raise ValueError("Refused write outside Shiur Library")
    data = value.encode("utf-8") if isinstance(value, str) else value
    mode = dropbox.files.WriteMode.overwrite if overwrite else dropbox.files.WriteMode.add
    try:
        dbx.files_upload(data, path, mode=mode, autorename=False, mute=True)
        return "saved"
    except Exception as ex:
        if not overwrite and "conflict" in str(ex).lower():
            return "exists"
        raise


def ensure(dbx, path):
    try:
        dbx.files_create_folder_v2(path)
    except Exception as ex:
        if "conflict" not in str(ex).lower():
            raise


def speaker_path(sid):
    if sid == "UNASSIGNED":
        return ROOT + "/Unassigned"
    return ROOT + "/Speakers/" + sid + " - " + SPEAKERS[sid]["name"]


def load_catalog(dbx):
    raw = read_bytes(dbx, ROOT + "/Indexes/master_catalog.json")
    if raw:
        result = json.loads(raw.decode("utf-8"))
    else:
        raw = read_bytes(dbx, ROOT + "/Indexes/master_catalog.csv")
        if not raw:
            raise RuntimeError("No existing Dropbox master catalog. Refusing to reset identifiers.")
        result = list(csv.DictReader(io.StringIO(raw.decode("utf-8-sig"))))
    for row in result:
        for key in ("shiur_id", "speaker_id", "video_id", "title", "youtube_url",
                    "transcript_status", "summary_status", "verification_status"):
            row.setdefault(key, "")
    return result


def normalized_words(value):
    """Normalize punctuation and HTML entities without loosening name boundaries."""
    value = unicodedata.normalize("NFKC", unescape(value or "")).casefold()
    return " ".join(re.sub(r"[^\w]+", " ", value).split())


def speaker_match(title, description):
    """Title attribution first; description requires explicit speaker wording."""
    title_words = normalized_words(title)
    title_hits = {sid for sid, spec in SPEAKERS.items()
                  if any(re.search(r"(?:^| )" + re.escape(normalized_words(alias)) +
                                   r"(?: |$)", title_words)
                         for alias in spec["aliases"])}
    if len(title_hits) == 1:
        return next(iter(title_hits))
    if title_hits:
        return "UNASSIGNED"  # Multiple speakers in title; human review.

    description_words = normalized_words(description)
    cues = r"(?:speaker|lecturer|taught by|shiur by|given by|presented by|with)"
    desc_hits = {sid for sid, spec in SPEAKERS.items()
                 if any(re.search(r"\b" + cues + r"\s+(?:rabbi\s+|rav\s+|reb\s+|r\s+)?" +
                                  re.escape(normalized_words(alias)) + r"\b",
                                  description_words)
                        for alias in spec["aliases"])}
    return next(iter(desc_hits)) if len(desc_hits) == 1 else "UNASSIGNED"


def reconcile_speakers(catalog, log):
    """Promote safely identified older leads, preserving established shiur IDs."""
    for entry in catalog:
        if "&#" in entry.get("title", "") or "&amp;" in entry.get("title", ""):
            entry["title"] = unescape(entry["title"])
        if entry.get("speaker_id") != "UNASSIGNED" or entry.get("shiur_id"):
            continue
        if entry.get("possible_duplicate_of"):
            continue
        # Deliberately rely on explicit names in titles for old unassigned leads.
        sid = speaker_match(entry.get("title", ""), "")
        if sid == "UNASSIGNED":
            continue
        entry["speaker_id"] = sid
        entry["verification_status"] = "metadata_found"
        entry["notes"] = ((entry.get("notes") or "") +
                          " Auto-attributed from explicit speaker name in title; provisional.").strip()
        log.append("ATTRIBUTED [" + sid + "] " + entry.get("video_id", "") +
                   " " + entry.get("title", ""))


def relevant(title, description):
    text = (title + " " + description[:800]).lower()
    subject = any(k in text for k in (
        "torah alef", "torah aleph", "torah 1", "torah #1",
        "torah א", "ashrei temimei", "true chen"))
    series = any(k in text for k in ("likutei moharan", "likutey moharan",
                                     "torah alef", "torah aleph"))
    return subject and series


def discover(catalog, log):
    if not YOUTUBE_KEY:
        log.append("YouTube API key absent; retained existing catalog.")
        return
    known = {x["video_id"] for x in catalog}
    since = (NOW - dt.timedelta(days=21)).strftime("%Y-%m-%dT%H:%M:%SZ")
    queries = ["Kollel Toras Chaim Likutei Moharan Torah 1",
               "Likutei Moharan Torah Aleph chabura"]
    queries += [spec["aliases"][0] + " Likutei Moharan Torah 1"
                for spec in SPEAKERS.values()]
    for query in queries:
        try:
            r = HTTP.get("https://www.googleapis.com/youtube/v3/search",
                params={"part": "snippet", "q": query, "type": "video",
                        "order": "date", "maxResults": 30,
                        "publishedAfter": since, "key": YOUTUBE_KEY}, timeout=25)
            r.raise_for_status()
            candidates = r.json().get("items", [])
        except Exception as ex:
            log.append("Discovery error: " + query + ": " + str(ex)[:180])
            continue
        for video in candidates:
            vid = video.get("id", {}).get("videoId", "")
            if not re.fullmatch(r"[A-Za-z0-9_-]{11}", vid) or vid in known:
                continue
            info = video.get("snippet", {})
            title = unescape(info.get("title", ""))
            description = info.get("description", "")
            if not relevant(title, description):
                continue
            sid = speaker_match(title, description)
            catalog.append(dict(speaker_id=sid, shiur_id="", legacy_id="",
                video_id=vid, title=title,
                youtube_url="https://www.youtube.com/watch?v=" + vid,
                posted_at=info.get("publishedAt", ""),
                channel=info.get("channelTitle", ""),
                transcript_status="not_extracted", summary_status="pending",
                verification_status=("attribution_unconfirmed" if sid == "UNASSIGNED"
                                     else "metadata_found"),
                possible_duplicate_of="", notes="New YouTube search result; speaker metadata provisional."))
            known.add(vid)
            log.append("NEW [" + sid + "] " + vid + " " + title)


def assign_ids(catalog):
    sequence = defaultdict(int)
    for entry in catalog:
        m = re.fullmatch(r"([A-Z]{3})-(\d{3})", entry.get("shiur_id") or "")
        if m:
            sequence[m.group(1)] = max(sequence[m.group(1)], int(m.group(2)))
    for entry in catalog:
        sid = entry["speaker_id"]
        if sid not in SPEAKERS or entry.get("shiur_id") or entry.get("possible_duplicate_of"):
            continue
        if entry.get("verification_status") not in (
            "verified_metadata", "metadata_found", "transcript_extracted",
            "metadata_verified_by_human"):
            continue
        sequence[sid] += 1
        entry["shiur_id"] = f"{sid}-{sequence[sid]:03d}"


def captions(video_id):
    """Return original caption segments without language reconstruction."""
    errors = []
    try:
        r = HTTP.get("https://yt-transcript.net/" + video_id, timeout=45)
        r.raise_for_status()
        value = r.json()
        if isinstance(value, list) and value and all(
                "start" in x and "text" in x for x in value):
            return value, "yt-transcript.net"
        errors.append("Transcript service returned no caption array")
    except Exception as ex:
        errors.append("Direct transcript: " + str(ex)[:150])
    try:
        from youtube_transcript_api import YouTubeTranscriptApi
        value = YouTubeTranscriptApi().fetch(video_id, languages=["en"])
        return [{"start": x.start, "duration": x.duration, "text": x.text}
                for x in value], "youtube-transcript-api"
    except Exception as ex:
        errors.append("YouTube fallback: " + str(ex)[:150])
    raise RuntimeError("; ".join(errors))


def timecode(seconds):
    n = int(float(seconds))
    return f"{n//3600:02d}:{(n//60)%60:02d}:{n%60:02d}"


def timed_text(segments):
    """45-second blocks. Keep original auto-caption wording, including errors."""
    blocks = []
    bucket = None
    current = []
    for item in segments:
        n = int(float(item["start"]) // 45)
        if bucket is not None and n != bucket:
            blocks.append("[" + timecode(bucket * 45) + "] " + " ".join(current))
            current = []
        bucket = n
        current.append(str(item["text"]).strip())
    if current:
        blocks.append("[" + timecode(bucket * 45) + "] " + " ".join(current))
    return "\n\n".join(blocks)


def word_file(heading, subtitle, sections, url=None):
    doc = Document()
    doc.styles["Normal"].font.size = Pt(9)
    section = doc.sections[0]
    section.top_margin = Inches(.7)
    section.bottom_margin = Inches(.7)
    doc.add_heading(heading, 0)
    doc.add_paragraph(subtitle)
    if url:
        doc.add_paragraph("Source: " + url)
    for header, content in sections:
        doc.add_heading(header, 1)
        for part in str(content).split("\n\n"):
            if part.strip():
                doc.add_paragraph(part.strip())
    data = io.BytesIO()
    doc.save(data)
    return data.getvalue()


def summary_from_captions(entry, text):
    if not OPENAI_KEY:
        return None
    from openai import OpenAI
    prompt = (
        "Write a careful Torah shiur study summary from AUTOMATIC CAPTIONS ONLY. "
        "They frequently mishear Hebrew and Yiddish. Do not invent quotations, "
        "pesukim, chazal sources, anecdotes, or speaker claims. Identify core "
        "ideas with approximate timestamps; mark ambiguous quotations [audio review]. "
        "Separate speaker's own teaching from inference. Use sections: Overview; "
        "Teachings with timestamps; Mekoros (verified vs tentative); "
        "Practical application; Questions; Caption uncertainties. "
        "This is a draft for human review.\n"
        f"Shiur ID: {entry['shiur_id']}\nTitle: {entry['title']}\n"
        f"Video: {entry['youtube_url']}\nTranscript:\n{text[:85000]}"
    )
    client = OpenAI(api_key=OPENAI_KEY)
    answer = client.responses.create(
        model=os.getenv("OPENAI_MODEL") or "gpt-4.1-mini",
        input=prompt, max_output_tokens=2300)
    return answer.output_text


def process_videos(dbx, catalog, log):
    for entry in catalog:
        sid = entry.get("speaker_id", "")
        ident = entry.get("shiur_id", "")
        if sid not in SPEAKERS or not ident:
            continue
        folder = speaker_path(sid)
        for sub in ("Raw", "Clean", "Summaries"):
            ensure(dbx, folder + "/" + sub)
        raw_path = folder + "/Raw/" + ident + " - captions.json"
        prior = read_bytes(dbx, raw_path)
        if prior:
            try:
                segments = json.loads(prior.decode("utf-8"))["segments"]
                entry["transcript_status"] = "raw_and_clean_archived"
            except Exception as ex:
                entry["transcript_status"] = "raw_unreadable"
                log.append(ident + " RAW UNREADABLE: " + str(ex)[:120])
                continue
        else:
            try:
                segments, source = captions(entry["video_id"])
                if not segments:
                    raise RuntimeError("Empty caption result")
                raw = {"source": entry["youtube_url"], "provider": source,
                       "retrieved_at": NOW.isoformat(), "language": "English/auto-unverified",
                       "segments": segments}
                save(dbx, raw_path, json.dumps(raw, ensure_ascii=False, indent=2))
                save(dbx, folder + "/Raw/" + ident + " - captions.txt",
                     timed_text(segments))
                save(dbx, folder + "/Clean/" + ident + " - Unverified Reading Edition.docx",
                     word_file(ident + " — Caption Reading Edition",
                        entry["title"], [
                            ("Editorial warning", "Automated captions. Hebrew/Yiddish and quoted sources "
                             "are unverified. Original words have not been reconstructed."),
                            ("Caption text", timed_text(segments))], entry["youtube_url"]))
                entry["transcript_status"] = "raw_and_clean_archived"
                entry["caption_provider"] = source
                entry["caption_segments"] = len(segments)
                log.append("ARCHIVED " + ident + " " + str(len(segments)) + " segments")
            except Exception as ex:
                entry["transcript_status"] = "extraction_or_archive_failed"
                log.append("EXTRACTION FAILED " + ident + ": " + str(ex)[:240])
                continue
        # User-initiated, source-aware commentary takes priority over unattended AI notes.
        manual_path = folder + "/Summaries/" + ident + " - Analytical Commentary v1.md"
        if read_bytes(dbx, manual_path):
            entry["summary_status"] = "manual_analytical_draft_needs_audio_review"
            entry["editorial_summary_path"] = manual_path
            continue
        summary_path = folder + "/Summaries/" + ident + " - Summary.md"
        if read_bytes(dbx, summary_path):
            entry["summary_status"] = "draft_from_auto_captions"
            continue
        if not OPENAI_KEY:
            entry["summary_status"] = "pending_no_openai_key"
            continue
        try:
            result = summary_from_captions(entry, timed_text(segments))
            if result:
                markdown = "# " + ident + " — " + entry["title"] + "\n\n" + \
                    "Draft from unverified captions. Source: " + entry["youtube_url"] + "\n\n" + result
                save(dbx, summary_path, markdown)
                save(dbx, folder + "/Summaries/" + ident + " - Summary.docx",
                    word_file(ident + " — Shiur Summary", entry["title"], [
                        ("Draft from unverified captions", result)], entry["youtube_url"]))
                entry["summary_status"] = "draft_from_auto_captions"
                log.append("SUMMARY " + ident + " saved")
        except Exception as ex:
            entry["summary_status"] = "summary_failed"
            log.append("SUMMARY FAILED " + ident + ": " + str(ex)[:210])


def build_indexes(dbx, catalog):
    keys = []
    for entry in catalog:
        for k in entry:
            if k not in keys:
                keys.append(k)
    output = io.StringIO()
    writer = csv.DictWriter(output, fieldnames=keys, extrasaction="ignore")
    writer.writeheader()
    writer.writerows(catalog)
    save(dbx, ROOT + "/Indexes/master_catalog.csv",
         "\ufeff" + output.getvalue(), overwrite=True)
    save(dbx, ROOT + "/Indexes/master_catalog.json",
         json.dumps(catalog, ensure_ascii=False, indent=2), overwrite=True)
    sections = []
    for sid, spec in SPEAKERS.items():
        rows = [x for x in catalog if x.get("speaker_id") == sid]
        text = "\n\n".join(
            (x.get("shiur_id") or "LEAD") + "  " + x.get("title", "") +
            "\n" + x.get("youtube_url", "") +
            "\nCaptions: " + x.get("transcript_status", "") +
            "; Summary: " + x.get("summary_status", "") +
            "; Verification: " + x.get("verification_status", "")
            for x in rows)
        sections.append((sid + " — " + spec["name"], text or "No indexed shiurim."))
        cumulative = []
        for e in rows:
            if not e.get("shiur_id"):
                continue
            # Prefer richer manual commentary when available; preserve pending otherwise.
            note = read_bytes(dbx, speaker_path(sid) + "/Summaries/" +
                              e["shiur_id"] + " - Analytical Commentary v1.md")
            if not note:
                note = read_bytes(dbx, speaker_path(sid) + "/Summaries/" +
                                  e["shiur_id"] + " - Summary.md")
            cumulative.append((e["shiur_id"] + " — " + e["title"],
                (note.decode("utf-8") if note else
                 "Pending transcript-backed summary.") +
                "\nSource: " + e["youtube_url"]))
        save(dbx, speaker_path(sid) + "/" + sid + " - Cumulative Summary.docx",
            word_file(spec["name"] + " — Cumulative Summary",
                "Automated source-backed archive; drafts require review.", cumulative),
            overwrite=True)
    other = [x for x in catalog if x.get("speaker_id") == "UNASSIGNED"]
    sections.append(("Unassigned — Needs verification", "\n\n".join(
        x.get("title", "") + "\n" + x.get("youtube_url", "")
        for x in other) or "None"))
    save(dbx, ROOT + "/Indexes/Master Shiur Index.docx",
         word_file("Torah Aleph — Master Shiur Index",
             "Updated " + NOW.strftime("%Y-%m-%d %H:%M UTC") +
             ". Speaker IDs and YouTube video IDs are stable.", sections),
         overwrite=True)


def main():
    dbx = dropbox_client()
    for path in (ROOT, ROOT + "/Indexes", ROOT + "/Speakers",
                 ROOT + "/Unassigned", ROOT + "/Comparative Learning", ROOT + "/Logs"):
        ensure(dbx, path)
    for sid in SPEAKERS:
        ensure(dbx, speaker_path(sid))
    catalog = load_catalog(dbx)
    log = ["Run " + NOW.isoformat(), "Baseline entries: " + str(len(catalog))]
    discover(catalog, log)
    reconcile_speakers(catalog, log)
    assign_ids(catalog)
    process_videos(dbx, catalog, log)
    build_indexes(dbx, catalog)
    log.append("Final catalog entries: " + str(len(catalog)))
    report = "\n".join(log) + "\n"
    save(dbx, ROOT + "/Logs/" + NOW.strftime("%Y-%m-%d") +
         " Daily Run.txt", report, overwrite=True)
    print(report)


if __name__ == "__main__":
    main()
